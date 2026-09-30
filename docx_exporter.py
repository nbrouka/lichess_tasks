"""
DOCX exporter for Lichess puzzles.
"""

import logging
import os
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Cm, Emu, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

from typing import Any, Optional, Sequence

from board_renderer import render_puzzle
from constants import (
    DOCX_MARGIN_CM,
    DOCX_PUZZLES_PER_PAGE,
    DOCX_COL_WIDTH_CM,
    DOCX_ROW_HEIGHT_TWIPS,
    DOCX_IMAGE_WIDTH_EMU,
    THUMBNAIL_SQUARE_SIZE,
    current_player_color_name,
    DOCX_TABLE_ROWS,
    DOCX_TABLE_COLS,
    DOCX_HEADER_LEFT_COL_WIDTH_CM,
    DOCX_HEADER_RIGHT_COL_WIDTH_CM,
    DOCX_CELL_MARGIN_TOP_DXA,
    DOCX_CELL_MARGIN_BOTTOM_DXA,
    DOCX_CELL_MARGIN_LEFT_DXA,
    DOCX_CELL_MARGIN_RIGHT_DXA,
    DOCX_CELL_BORDER_VAL,
    DOCX_CELL_BORDER_SZ,
    DOCX_PARAGRAPH_SPACE_BEFORE_PT,
    DOCX_CAPTION_SPACE_AFTER_PT,
    DOCX_PARAGRAPH_LINE_SPACING_PT,
)

logger = logging.getLogger(__name__)


def _set_cell_border(cell, **kwargs: Any) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), str(edge_data.get('val', DOCX_CELL_BORDER_VAL)))
            element.set(qn('w:sz'), str(edge_data.get('sz', DOCX_CELL_BORDER_SZ)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), str(edge_data.get('color', 'auto')))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def _set_row_height(row, height_twips: int) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(height_twips))
    trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)


def _set_cell_margins(cell, top: Optional[int] = None, bottom: Optional[int] = None, left: Optional[int] = None, right: Optional[int] = None) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    margins = {
        'top': top if top is not None else DOCX_CELL_MARGIN_TOP_DXA,
        'bottom': bottom if bottom is not None else DOCX_CELL_MARGIN_BOTTOM_DXA,
        'left': left if left is not None else DOCX_CELL_MARGIN_LEFT_DXA,
        'right': right if right is not None else DOCX_CELL_MARGIN_RIGHT_DXA,
    }
    for edge, val in margins.items():
        elem = OxmlElement('w:{}'.format(edge))
        elem.set(qn('w:w'), str(val))
        elem.set(qn('w:type'), 'dxa')
        tcMar.append(elem)
    tcPr.append(tcMar)


class PuzzleDocxExporter:
    """Exports selected puzzles to a DOCX document."""

    def __init__(self, puzzles: Sequence[Any], topic: str) -> None:
        self.puzzles = puzzles
        self.topic = topic

    def export(self, path: str) -> None:
        logger.info("DOCX export started path=%s topic=%s total_puzzles=%d", path, self.topic, len(self.puzzles))
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Cm(DOCX_MARGIN_CM)
        section.bottom_margin = Cm(DOCX_MARGIN_CM)
        section.left_margin = Cm(DOCX_MARGIN_CM)
        section.right_margin = Cm(DOCX_MARGIN_CM)

        temp_files = []
        try:
            for page_idx in range(0, len(self.puzzles), DOCX_PUZZLES_PER_PAGE):
                page_puzzles = self.puzzles[page_idx:page_idx + DOCX_PUZZLES_PER_PAGE]
                sheet_num = page_idx // DOCX_PUZZLES_PER_PAGE + 1
                logger.info("Processing page %d puzzles=%d", sheet_num, len(page_puzzles))

                if page_idx > 0:
                    doc.add_page_break()

                header_table = doc.add_table(rows=1, cols=2)
                header_table.autofit = False
                header_table.columns[0].width = Cm(DOCX_HEADER_LEFT_COL_WIDTH_CM)
                header_table.columns[1].width = Cm(DOCX_HEADER_RIGHT_COL_WIDTH_CM)

                cell_left = header_table.cell(0, 0)
                cell_left.text = f"Лист {sheet_num}"
                cell_left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                cell_left.paragraphs[0].paragraph_format.space_after = Pt(0)

                cell_center = header_table.cell(0, 1)
                cell_center.text = self.topic
                cell_center.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cell_center.paragraphs[0].paragraph_format.space_after = Pt(0)

                for row in header_table.rows:
                    for cell in row.cells:
                        _set_cell_border(cell, top={"val": DOCX_CELL_BORDER_VAL, "sz": DOCX_CELL_BORDER_SZ}, bottom={"val": DOCX_CELL_BORDER_VAL, "sz": DOCX_CELL_BORDER_SZ}, left={"val": DOCX_CELL_BORDER_VAL, "sz": DOCX_CELL_BORDER_SZ}, right={"val": DOCX_CELL_BORDER_VAL, "sz": DOCX_CELL_BORDER_SZ})

                table = doc.add_table(rows=DOCX_TABLE_ROWS, cols=DOCX_TABLE_COLS)
                table.autofit = False
                table.allow_autofit = False
                table.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for col in table.columns:
                    col.width = Cm(DOCX_COL_WIDTH_CM)

                for row in table.rows:
                    _set_row_height(row, DOCX_ROW_HEIGHT_TWIPS)
                    for cell in row.cells:
                        _set_cell_margins(
                            cell,
                            top=DOCX_CELL_MARGIN_TOP_DXA,
                            bottom=DOCX_CELL_MARGIN_BOTTOM_DXA,
                            left=DOCX_CELL_MARGIN_LEFT_DXA,
                            right=DOCX_CELL_MARGIN_RIGHT_DXA,
                        )
                        _set_cell_border(cell, top={"val": DOCX_CELL_BORDER_VAL, "sz": DOCX_CELL_BORDER_SZ}, bottom={"val": DOCX_CELL_BORDER_VAL, "sz": DOCX_CELL_BORDER_SZ}, left={"val": DOCX_CELL_BORDER_VAL, "sz": DOCX_CELL_BORDER_SZ}, right={"val": DOCX_CELL_BORDER_VAL, "sz": DOCX_CELL_BORDER_SZ})

                for i, puzzle in enumerate(page_puzzles):
                    row_idx = i // 3
                    col_idx = i % 3
                    cell = table.cell(row_idx, col_idx)

                    global_idx = page_idx + i + 1
                    logger.debug("Rendering puzzle %d/%d puzzle_id=%s", global_idx, len(self.puzzles), puzzle.puzzle_id)
                    img = render_puzzle(puzzle.fen, puzzle.moves, square_size=THUMBNAIL_SQUARE_SIZE, move_index=1, show_coordinates=False)
                    pixels = img.load()
                    original_pixel = pixels[0, 0]
                    pixels[0, 0] = (
                        original_pixel[0],
                        original_pixel[1],
                        (original_pixel[2] + global_idx) % 256,
                    )
                    fd, img_path = tempfile.mkstemp(suffix=".png")
                    os.close(fd)
                    img.save(img_path)
                    temp_files.append(img_path)
                    logger.debug("Saved temporary image %s size=%s", img_path, img.size)

                    p0 = cell.paragraphs[0]
                    p0.paragraph_format.space_after = Pt(0)
                    run = p0.add_run()
                    run.add_picture(img_path, width=Emu(DOCX_IMAGE_WIDTH_EMU))
                    run.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    run.font.size = Pt(9)
                    run.font.name = 'Cambria'

                    p = cell.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(DOCX_PARAGRAPH_SPACE_BEFORE_PT)
                    p.paragraph_format.space_after = Pt(DOCX_CAPTION_SPACE_AFTER_PT)
                    p.paragraph_format.line_spacing = Pt(DOCX_PARAGRAPH_LINE_SPACING_PT)
                    color_text = current_player_color_name(puzzle.color)
                    p.add_run(f"№{global_idx}. {color_text}")

            logger.info("Saving document path=%s", path)
            doc.save(path)
            logger.info("DOCX saved successfully path=%s", path)
        finally:
            for p in temp_files:
                try:
                    os.unlink(p)
                except OSError:
                    pass
