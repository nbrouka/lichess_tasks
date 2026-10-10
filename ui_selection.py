"""
Mixin for selected puzzles panel in Lichess Puzzle Viewer.
"""

import logging
import tkinter as tk
from pathlib import Path
from tkinter import ttk, messagebox, filedialog
from PIL import ImageTk

from board_renderer import render_puzzle
from constants import t, THUMBNAIL_SQUARE_SIZE, DOCX_DIALOG_GEOMETRY, current_player_color_name, DOCX_ANSWER_MARGIN_CM, DOCX_ANSWER_SPACE_AFTER_PT, DELETE_BUTTON_WIDTH, DELETE_BUTTON_HEIGHT
from database import Puzzle
from docx_exporter import PuzzleDocxExporter
from pdf_exporter import PuzzlePdfExporter
from docx import Document
from docx.shared import Cm, Pt

LOG_PATH = Path(__file__).parent / "docx_export.log"
logger = logging.getLogger("docx_export")
if not logger.handlers:
    logger.setLevel(logging.DEBUG)
    _handler = logging.FileHandler(LOG_PATH, encoding="utf-8")
    _handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(_handler)


class PuzzleSelectionMixin:
    def _create_selected_widget(self, puzzle: Puzzle) -> None:
        square_size = THUMBNAIL_SQUARE_SIZE
        thumb = render_puzzle(puzzle.fen, puzzle.moves, square_size=square_size, move_index=1, show_coordinates=False)
        photo = ImageTk.PhotoImage(thumb)

        row = len(self.selected_puzzles) - 1
        wrapper = tk.Frame(self.selected_inner, borderwidth=0, padx=0, pady=0)
        wrapper.grid(row=row, column=0, sticky="nw", padx=4, pady=4)

        img_label = tk.Label(wrapper, image=photo, borderwidth=0)
        img_label.image = photo
        img_label.pack(side=tk.TOP)

        del_btn = tk.Button(
            wrapper,
            text="✕",
            width=DELETE_BUTTON_WIDTH,
            height=DELETE_BUTTON_HEIGHT,
            command=lambda p=puzzle, w=wrapper: self._remove_from_selected(p, w),
        )
        del_btn.pack(side=tk.TOP, fill=tk.X)

        self.selected_inner.grid_rowconfigure(row, weight=0, minsize=thumb.size[1] + 20)
        self.selected_inner.grid_columnconfigure(0, weight=0, minsize=thumb.size[0])

    def _add_to_selected(self) -> None:
        if self.current_index is None or not self.filtered_puzzles:
            return
        puzzle = self.filtered_puzzles[self.current_index]

        if any(p.fen == puzzle.fen for p in self.selected_puzzles):
            return

        self.selected_puzzles.append(puzzle)
        self._create_selected_widget(puzzle)
        self.selected_inner.update_idletasks()
        self.selected_canvas.configure(scrollregion=self.selected_canvas.bbox("all"))
        self._update_selected_title()
        self._save_session()

    def _update_selected_title(self) -> None:
        self.selected_title_label.config(text=f"{t('selected_title')} ({len(self.selected_puzzles)})")

    def _on_selected_canvas_resize(self, event: tk.Event) -> None:
        """
        Обновляет ширину внутреннего окна при изменении размера канваса.

        Нужно, чтобы скроллбар работал корректно при изменении размера окна.
        """
        self.selected_canvas.itemconfig(self._selected_window_id, width=event.width)

    def _add_to_theme_dialog(self) -> None:
        """
        Открывает диалог добавления выбранных задач в тему.

        Пользователь может выбрать существующую тему или создать новую.
        Задачи привязываются к теме в БД, файлы не создаются.
        """
        logger.info("Open add to theme dialog, selected puzzles count=%d", len(self.selected_puzzles))
        if not self.selected_puzzles:
            messagebox.showinfo(t("about_title"), t("msg_no_selected_puzzles"))
            return

        user_themes = self.db.get_user_themes()

        dialog = tk.Toplevel(self.root)
        dialog.title("Добавить в тему")
        dialog.geometry(DOCX_DIALOG_GEOMETRY)
        dialog.transient(self.root)
        dialog.grab_set()

        ttk.Label(dialog, text="Существующая тема:").pack(anchor=tk.W, padx=10, pady=(10, 4))
        theme_var = tk.StringVar(value="")
        theme_cb = ttk.Combobox(
            dialog, textvariable=theme_var,
            values=[""] + user_themes, state="readonly", width=36,
        )
        theme_cb.pack(anchor=tk.W, padx=10, pady=(0, 10))

        ttk.Label(dialog, text="Или создать новую:").pack(anchor=tk.W, padx=10)
        new_theme_var = tk.StringVar(value="")
        new_theme_entry = ttk.Entry(dialog, textvariable=new_theme_var)
        new_theme_entry.pack(fill=tk.X, padx=10, pady=(0, 10))
        new_theme_entry.focus_set()

        def on_add() -> None:
            theme_name = new_theme_var.get().strip() or theme_var.get()
            if not theme_name:
                messagebox.showwarning("Внимание", "Укажите тему.", parent=dialog)
                return
            try:
                theme_id = self.db.get_or_create_user_theme(theme_name)
                puzzle_ids = [p.puzzle_id for p in self.selected_puzzles]
                self.db.link_puzzles_to_user_theme(theme_id, puzzle_ids)
                self._refresh_user_themes()
                self._refresh_exclude_user_themes()
                self.status_label.config(text=f"Задачи добавлены в тему «{theme_name}»")
                logger.info("Puzzles added to theme name=%s count=%d", theme_name, len(puzzle_ids))
                dialog.destroy()
                self._clear_selection()
            except Exception as exc:
                logger.exception("Failed to add puzzles to theme name=%s", theme_name)
                messagebox.showerror("Ошибка", f"Не удалось добавить задачи в тему: {exc}", parent=dialog)

        ttk.Button(dialog, text="Добавить", command=on_add).pack(pady=(0, 10))

    def _remove_from_selected(self, puzzle: Puzzle, wrapper: tk.Frame) -> None:
        if puzzle in self.selected_puzzles:
            self.selected_puzzles.remove(puzzle)
        wrapper.destroy()
        self.selected_inner.update_idletasks()
        self.selected_canvas.configure(scrollregion=self.selected_canvas.bbox("all"))
        self._update_selected_title()
        self._save_session()

    def _create_sheets_docx(self, path: str, topic: str) -> None:
        """
        Creates DOCX with puzzles and separate DOCX with answers.

        Also saves topic to user themes DB if specified.
        """
        puzzles = list(self.selected_puzzles)
        exporter = PuzzleDocxExporter(puzzles, topic)
        exporter.export(path)

        answers_path = path.replace(".docx", "_ответы.docx")
        self._create_answers_docx(answers_path, topic, puzzles)

        if topic:
            self._save_topic_to_db(topic)
            self.db.mark_theme_sheets_created(topic)

    def _create_sheets_pdf(self, path: str, topic: str) -> None:
        puzzles = list(self.selected_puzzles)
        exporter = PuzzlePdfExporter(puzzles, topic)
        exporter.export_with_answers(path)

        if topic:
            self._save_topic_to_db(topic)
            self.db.mark_theme_sheets_created(topic)

    def _create_sheets_by_theme_dialog(self) -> None:
        """
        Открывает диалог создания листов по существующей теме (из меню).

        Показывает темы с пометкой "(уже созданы)" для тем, по которым уже
        создавались листы (PDF/DOCX).
        Позволяет выбрать тему и создать PDF/DOCX листы.
        """
        logger.info("Open create sheets by theme dialog")
        user_themes = self.db.get_user_themes()
        if not user_themes:
            messagebox.showinfo("Создать листы", "Нет пользовательских тем. Сначала добавьте задачи в тему.")
            return

        dialog = tk.Toplevel(self.root)
        dialog.title("Создать листы")
        dialog.geometry(DOCX_DIALOG_GEOMETRY)
        dialog.transient(self.root)
        dialog.grab_set()

        ttk.Label(dialog, text="Тема:").pack(anchor=tk.W, padx=10, pady=(10, 4))

        # Пометка «уже созданы» — только для тем, по которым реально создавались листы.
        themes_with_sheets = set(self.db.get_themes_with_sheets_created())
        theme_display_map = {}
        display_values = []
        for name in user_themes:
            display = f"{name} (уже созданы)" if name in themes_with_sheets else name
            theme_display_map[display] = name
            display_values.append(display)

        theme_var = tk.StringVar(value="")
        theme_cb = ttk.Combobox(
            dialog, textvariable=theme_var,
            values=display_values, state="readonly", width=36,
        )
        theme_cb.pack(anchor=tk.W, padx=10, pady=(0, 10))
        theme_cb.focus_set()

        def on_create() -> None:
            display_name = theme_var.get()
            if not display_name:
                messagebox.showwarning("Внимание", "Выберите тему.", parent=dialog)
                return
            theme_name = theme_display_map.get(display_name, display_name)
            try:
                # Получаем задачи, привязанные к этой теме
                cursor = self.db.conn.cursor()
                cursor.execute(
                    """
                    SELECT p.* FROM puzzles p
                    JOIN puzzle_user_themes put ON p.PuzzleId = put.puzzle_id
                    JOIN user_themes ut ON put.theme_id = ut.id
                    WHERE ut.name = ?
                    """,
                    (theme_name,),
                )
                rows = cursor.fetchall()
                if not rows:
                    messagebox.showinfo("Создать листы", f"В теме «{theme_name}» нет задач.", parent=dialog)
                    return

                puzzles = [self.db._row_to_puzzle(row) for row in rows]
                logger.info("Creating sheets for theme=%s puzzles=%d", theme_name, len(puzzles))

                # Выбор пути сохранения
                path = filedialog.asksaveasfilename(
                    defaultextension=".pdf",
                    filetypes=[("PDF files", "*.pdf"), ("DOCX files", "*.docx")],
                    initialfile=f"{theme_name}.pdf",
                    parent=dialog,
                )
                if not path:
                    return

                if path.lower().endswith(".docx"):
                    exporter = PuzzleDocxExporter(puzzles, theme_name)
                    exporter.export(path)
                    answers_path = path.replace(".docx", "_ответы.docx")
                    self._create_answers_docx(answers_path, theme_name, puzzles)
                else:
                    exporter = PuzzlePdfExporter(puzzles, theme_name)
                    exporter.export_with_answers(path)

                self.db.mark_theme_sheets_created(theme_name)
                self._refresh_user_themes()
                self._refresh_exclude_user_themes()
                messagebox.showinfo("Готово", f"Листы сохранены: {path}", parent=dialog)
                dialog.destroy()
            except Exception as exc:
                logger.exception("Failed to create sheets for theme=%s", theme_name)
                messagebox.showerror("Ошибка", f"Не удалось создать листы: {exc}", parent=dialog)

        ttk.Button(dialog, text="Создать", command=on_create).pack(pady=(0, 10))

    def _create_answers_docx(self, path: str, topic: str, puzzles: list) -> None:
        """
        Создаёт DOCX с ответами для всех задач.

        В файле указывается номер задачи, цвет хода и форматированное решение.
        """
        logger.info("Creating answers DOCX path=%s puzzles=%d", path, len(puzzles))

        doc = Document()
        section = doc.sections[0]
        section.top_margin = Cm(DOCX_ANSWER_MARGIN_CM)
        section.bottom_margin = Cm(DOCX_ANSWER_MARGIN_CM)
        section.left_margin = Cm(DOCX_ANSWER_MARGIN_CM)
        section.right_margin = Cm(DOCX_ANSWER_MARGIN_CM)

        doc.add_paragraph(f"Тема: {topic}")

        for idx, puzzle in enumerate(puzzles, 1):
            formatted_solution = self._format_answers_solution(puzzle)
            color_text = current_player_color_name(puzzle.color)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(DOCX_ANSWER_SPACE_AFTER_PT)
            p.add_run(f"№{idx}. {color_text}: {formatted_solution}")

        doc.save(path)

    def _save_topic_to_db(self, topic: str) -> None:
        """
        Сохраняет тему в пользовательские темы и связывает с выбранными задачами.
        """
        theme_id = self.db.get_or_create_user_theme(topic)
        puzzle_ids = [p.puzzle_id for p in self.selected_puzzles]
        self.db.link_puzzles_to_user_theme(theme_id, puzzle_ids)
        self._refresh_user_themes()
        self._refresh_exclude_user_themes()

    def _format_answers_solution(self, puzzle: Puzzle) -> str:
        """
        Форматирует решение для файла ответов.

        Если черные ходят первыми, добавляет "..." после номера хода.
        """
        moves = puzzle.moves.split()
        if not moves:
            return ""

        is_black_first = puzzle.color == "w"
        parts = []
        move_number = 1
        i = 0

        if is_black_first:
            parts.append(f"{move_number}. ...{moves[i]}")
            i = 1
            move_number = 2

        while i < len(moves):
            if i + 1 < len(moves):
                parts.append(f"{move_number}.{moves[i]} {moves[i + 1]}")
                i += 2
            else:
                parts.append(f"{move_number}.{moves[i]}")
                i += 1
            move_number += 1

        return " ".join(parts)

    def _clear_selection(self) -> None:
        self.selected_puzzles.clear()
        for widget in self.selected_inner.winfo_children():
            widget.destroy()
        self.selected_inner.update_idletasks()
        self.selected_canvas.configure(scrollregion=self.selected_canvas.bbox("all"))
        self._update_selected_title()
        self._save_session()
